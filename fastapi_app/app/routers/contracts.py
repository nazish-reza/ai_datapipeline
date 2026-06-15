"""routers/contracts.py"""
import os, re, logging, tempfile
from pathlib import Path
import yaml
from fastapi import APIRouter, HTTPException, UploadFile, File
from fastapi.responses import JSONResponse
from pydantic import BaseModel

router = APIRouter()
CONTRACTS_DIR = Path(os.getenv("CONTRACTS_DIR", "contracts"))
logging.getLogger("datacontract.imports.excel_importer").setLevel(logging.ERROR)

def _parse_yaml(text):
    try: return yaml.safe_load(text) or {}
    except: return {}

def _get_columns(data):
    schema = data.get("schema", [])
    if not schema: return []
    t = schema[0] if isinstance(schema, list) else schema
    return t.get("properties", t.get("columns", []))

def _get_description(data):
    d = data.get("description", "")
    return d.get("purpose","") if isinstance(d,dict) else str(d) if d else ""

def _get_custom(data, key):
    for cp in data.get("customProperties", []):
        if isinstance(cp, dict) and cp.get("property") == key:
            v = cp.get("value", "")
            if isinstance(v, str) and v.startswith("="): return ""
            return str(v).strip() if v else ""
    return ""

def _get_owner(data):
    return next((t.get("username","") for t in data.get("team",[]) if t.get("role")=="owner"), "")

def _get_sources(data):
    src_str = _get_custom(data, "sources") or _get_custom(data, "upstream_entities")
    if src_str:
        return [s.strip() for s in src_str.split(",") if s.strip()]
    cols = _get_columns(data)
    sources = set()
    for c in cols:
        srcs = c.get("transformSourceObjects", c.get("transformSources", []))
        if isinstance(srcs, list):
            for s in srcs:
                if s and str(s).lower() not in ("derived","nan","n/a",""): sources.add(str(s).strip())
        elif isinstance(srcs, str) and srcs.strip() and srcs.lower() not in ("derived","nan","n/a",""):
            sources.add(srcs.strip())
    return sorted(sources)

def _get_data_product_type(data):
    for key in ["dataProductType", "data_product_type"]:
        val = _get_custom(data, key)
        if val:
            v = val.lower()
            if "consumer" in v: return "Data Product Consumer-Aligned"
            elif "source" in v or "aligned" in v: return "Data Product Source-Aligned"
            return val
    return ""

def _contract_meta(data):
    cols = _get_columns(data)
    return {
        "id": data.get("id",""), "version": data.get("version",""),
        "status": data.get("status",""), "domain": data.get("domain",""),
        "title": data.get("name", data.get("id","")),
        "description": _get_description(data), "owner": _get_owner(data),
        "tags": data.get("tags",[]), "column_count": len(cols),
        "has_quality": bool(data.get("quality")),
        "has_relations": bool(data.get("relationships")),
        "layer": _get_custom(data, "layer"),
        "dataProductType": _get_data_product_type(data),
        "sources": _get_sources(data),
    }

def _read_server_details_from_excel(excel_path):
    """Read project and dataset from the Excel Servers sheet (datacontract CLI misses these)."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(excel_path, data_only=True)
        if "Servers" not in wb.sheetnames:
            return {}
        ws = wb["Servers"]
        details = {}
        for r in range(1, ws.max_row + 2):
            for col in [1, 2]:
                key = str(ws.cell(row=r, column=col).value or "").strip().lower()
                if not key:
                    continue
                # Value is typically in column C (index 3) or column B+1
                val_col = 3 if col <= 2 else col + 1
                val = str(ws.cell(row=r, column=val_col).value or "").strip()
                if not val or val.startswith("="):
                    continue
                if "project" in key:
                    details["project"] = val
                elif "dataset" in key:
                    details["dataset"] = val
                elif "environment" in key and "environment" not in details:
                    details["environment"] = val
                elif "description" in key and "description" not in details:
                    details["description"] = val
        return details
    except Exception:
        return {}

def _enrich_yaml_with_server(yaml_text, server_details):
    """Inject project/dataset into the servers section of the YAML."""
    if not server_details:
        return yaml_text
    parsed = _parse_yaml(yaml_text)
    if not parsed:
        return yaml_text

    servers = parsed.get("servers", [])
    if servers:
        for s in servers:
            if isinstance(s, dict):
                if "project" not in s and server_details.get("project"):
                    s["project"] = server_details["project"]
                if "dataset" not in s and server_details.get("dataset"):
                    s["dataset"] = server_details["dataset"]
    else:
        # No servers at all — create one
        parsed["servers"] = [{
            "server": "production",
            "type": "BigQuery",
            "environment": server_details.get("environment", "production"),
            "project": server_details.get("project", ""),
            "dataset": server_details.get("dataset", ""),
        }]

    return yaml.dump(parsed, default_flow_style=False, allow_unicode=True, sort_keys=False, indent=2)

def _import_excel(file_path):
    try:
        from datacontract.data_contract import DataContract
        dc = DataContract()
        result = dc.import_from_source("excel", source=file_path)
        yaml_text = result.to_yaml()

        # Post-process: inject server details that the CLI misses
        server_details = _read_server_details_from_excel(file_path)
        if server_details:
            yaml_text = _enrich_yaml_with_server(yaml_text, server_details)

        return yaml_text
    except ImportError:
        raise HTTPException(500, "datacontract-cli[excel] not installed")
    except Exception as e:
        raise HTTPException(422, f"datacontract import failed: {str(e)}")

def _remove_old_versions(domain_dir, entity):
    removed = []
    for old in domain_dir.glob(f"{entity}_v*.yaml"):
        old.unlink(); removed.append(str(old))
    return removed

@router.post("/import")
async def import_contract(file: UploadFile = File(...)):
    if not file.filename.endswith(".xlsx"): raise HTTPException(400, "Only .xlsx")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(await file.read()); tmp_path = tmp.name
    try: yaml_text = _import_excel(tmp_path)
    finally:
        if os.path.exists(tmp_path): os.unlink(tmp_path)
    if not yaml_text or not yaml_text.strip(): raise HTTPException(422, "Empty YAML")
    parsed = _parse_yaml(yaml_text)
    return JSONResponse({"yaml":yaml_text,"parsed":_contract_meta(parsed),"domain":parsed.get("domain","default"),"filename":file.filename})

class SaveRequest(BaseModel):
    yaml_text: str; domain: str; entity: str; version: str = "1.0.0"

@router.post("/save")
async def save_contract(req: SaveRequest):
    parsed = _parse_yaml(req.yaml_text)
    if not parsed: raise HTTPException(400, "Invalid YAML")
    domain = re.sub(r"[^a-z0-9_\-]","_",req.domain.lower())
    entity = re.sub(r"[^a-z0-9_\-]","_",req.entity.lower())
    version = re.sub(r"[^0-9\.]","",req.version) or "1.0.0"
    out_dir = CONTRACTS_DIR / domain; out_dir.mkdir(parents=True, exist_ok=True)
    removed = _remove_old_versions(out_dir, entity)
    out_path = out_dir / f"{entity}_v{version}.yaml"
    out_path.write_text(req.yaml_text)
    msg = f"Saved to {out_path}"
    if removed: msg += f" (replaced {len(removed)} old version(s))"
    return JSONResponse({"saved":True,"path":str(out_path),"message":msg})

@router.get("")
async def list_contracts():
    if not CONTRACTS_DIR.exists(): return JSONResponse({"domains":[],"total":0})
    domains = {}; total = 0
    for yf in sorted(CONTRACTS_DIR.rglob("*.yaml")):
        try:
            data = _parse_yaml(yf.read_text()); meta = _contract_meta(data)
            domain = yf.parent.name
            meta["file"] = str(yf); meta["entity"] = yf.stem.split("_v")[0]
            if domain not in domains: domains[domain] = []
            domains[domain].append(meta); total += 1
        except: continue
    return JSONResponse({"domains":[{"domain":d,"contracts":c} for d,c in sorted(domains.items())],"total":total})

@router.get("/{domain}/{entity}")
async def get_contract(domain: str, entity: str):
    matches = sorted(CONTRACTS_DIR.glob(f"{domain}/{entity}_v*.yaml"), reverse=True)
    if not matches: raise HTTPException(404, f"Not found: {domain}/{entity}")
    yf = matches[0]
    try: text = yf.read_text(); parsed = _parse_yaml(text)
    except Exception as e: raise HTTPException(500, f"Read error: {e}")
    raw_cols = _get_columns(parsed)
    columns = [{
        "name":col.get("name",""), "logicalType":col.get("logicalType",""),
        "physicalType":col.get("physicalType",""), "description":col.get("description",""),
        "required":col.get("required",False), "unique":col.get("unique",False),
        "primaryKey":col.get("primaryKey",False), "businessName":col.get("businessName",""),
        "transformLogic":col.get("transformLogic",""),
        "transformSources":col.get("transformSourceObjects",col.get("transformSources",[])),
        "transformDescription":col.get("transformDescription",""),
        "examples":col.get("examples",[]), "classification":col.get("classification",""),
        "quality":col.get("quality",[]),
    } for col in raw_cols]
    return JSONResponse({
        "id":parsed.get("id",""), "apiVersion":parsed.get("apiVersion",""),
        "version":parsed.get("version",""), "status":parsed.get("status",""),
        "domain":parsed.get("domain",""), "dataProduct":parsed.get("dataProduct",""),
        "name":parsed.get("name",""),
        "layer": _get_custom(parsed, "layer"),
        "dataProductType": _get_data_product_type(parsed),
        "sources": _get_sources(parsed),
        "info":{"title":parsed.get("name",""),"description":_get_description(parsed),
                "owner":_get_owner(parsed),"purpose":_get_description(parsed),"tags":parsed.get("tags",[])},
        "columns":columns, "quality":parsed.get("quality",[]),
        "relationships":parsed.get("relationships",[]),
        "team":parsed.get("team",[]), "servers":parsed.get("servers",[]),
        "tags":parsed.get("tags",[]), "customProperties":parsed.get("customProperties",[]),
        "yaml_text":text, "file":str(yf),
    })